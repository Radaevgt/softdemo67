"""Отрисовка протокола в .docx — редактируемый файл для СЭДО и печати."""
from __future__ import annotations

import io

from docx import Document as DocxDocument
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from .layout import Document, Paragraph, Steps, Table

BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(11)
WARNING_COLOR = RGBColor(0xA3, 0x2B, 0x2B)
NOTE_COLOR = RGBColor(0x4A, 0x57, 0x69)
PAGE_WIDTH_CM = 16.5  # A4 минус поля


def _style(document: DocxDocument) -> None:
    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.paragraph_format.space_after = Pt(6)


def _paragraph(document: DocxDocument, block: Paragraph) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(block.text)
    if block.emphasis == "lead":
        run.bold = True
    elif block.emphasis == "warning":
        run.bold = True
        run.font.color.rgb = WARNING_COLOR
    elif block.emphasis == "note":
        run.italic = True
        run.font.color.rgb = NOTE_COLOR
        run.font.size = Pt(10)


def _table(document: DocxDocument, block: Table) -> None:
    columns = len(block.header) if block.header else max(len(row) for row in block.rows)
    table = document.add_table(rows=0, cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    if block.header:
        cells = table.add_row().cells
        for index, title in enumerate(block.header):
            cells[index].text = ""
            run = cells[index].paragraphs[0].add_run(title)
            run.bold = True

    for row in block.rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = ""
            paragraph = cells[index].paragraphs[0]
            run = paragraph.add_run(value)
            # В таблице «поле — значение» подпись поля выделяется.
            if not block.header and index == 0:
                run.bold = True

    if block.widths:
        from docx.shared import Cm

        for row in table.rows:
            for index, width in enumerate(block.widths[:columns]):
                row.cells[index].width = Cm(PAGE_WIDTH_CM * width)


def _steps(document: DocxDocument, block: Steps) -> None:
    number = 0
    for step in block.items:
        if step.kind == "header":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(step.text)
            run.bold = True
            run.italic = True
            continue
        number += 1
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Pt(18)
        paragraph.add_run(f"{number}. {step.text}")


def render(source: Document) -> bytes:
    document = DocxDocument()
    _style(document)

    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(source.title)
    run.bold = True
    run.font.size = Pt(13)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(source.subtitle)
    run.font.size = Pt(10)
    run.font.color.rgb = NOTE_COLOR

    for section in source.sections:
        if section.title:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(14)
            run = paragraph.add_run(section.title)
            run.bold = True
            run.font.size = Pt(12)

        for block in section.blocks:
            if isinstance(block, Paragraph):
                _paragraph(document, block)
            elif isinstance(block, Table):
                _table(document, block)
                document.add_paragraph()
            elif isinstance(block, Steps):
                _steps(document, block)

    if source.signature_lines:
        document.add_paragraph()
        for line in source.signature_lines:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(18)
            paragraph.add_run(line)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
